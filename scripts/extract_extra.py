# -*- coding: utf-8 -*-
"""
Trích xuất các nhóm dữ liệu bổ sung từ workbook lương:
CTV/TTS, biến động lương, truy thu, thưởng Ebitda, KPI chi tiết, thang bảng lương.
Cách dùng:  python scripts/extract_extra.py "duong-dan-file.xlsx" [password]
"""
import msoffcrypto, io, json, datetime, sys, os, re
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

SRC = sys.argv[1]
PASSWORD = sys.argv[2] if len(sys.argv) > 2 else "2019OZ"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")
MONTH = "2026-07"

def norm(v): return re.sub(r"\s+", " ", str(v)).strip().lower() if v is not None else ""
def num(v): return round(float(v), 2) if isinstance(v, (int, float)) else None
def s(v):
    t = str(v).strip() if v is not None else None
    return t if t else None
def d(v): return v.strftime("%Y-%m-%d") if isinstance(v, (datetime.datetime, datetime.date)) else None
def is_code(v):
    t = s(v)
    return bool(t) and len(t) >= 6 and not t.isdigit()

with open(SRC, "rb") as f:
    of = msoffcrypto.OfficeFile(f)
    of.load_key(password=PASSWORD)
    dec = io.BytesIO()
    of.decrypt(dec)
dec.seek(0)
wb = openpyxl.load_workbook(dec, data_only=True)

# ---------- 1. CTV (bảng lương cộng tác viên) ----------
ctv_employees, ctv_salaries = {}, []
for sheet, company in [("CTV XTRADING. (2)", "XTRADING"), ("CTV OZO", "OZO")]:
    if sheet not in wb.sheetnames: continue
    ws = wb[sheet]
    for r in range(4, ws.max_row + 1):
        code, name = s(ws.cell(r, 2).value), s(ws.cell(r, 3).value)
        if not is_code(code) or not name: continue
        g = lambda L: num(ws[f"{L}{r}"].value)
        status_txt = norm(ws[f"H{r}"].value)
        ctv_employees[code] = {
            "employee_code": code, "full_name": name,
            "position": s(ws[f"D{r}"].value), "department": s(ws[f"E{r}"].value),
            "company": company, "employee_type": "ctv",
            "join_date": d(ws[f"F{r}"].value), "contract_start_date": d(ws[f"G{r}"].value),
            "contract_type": s(ws[f"J{r}"].value) or "CTV",
            "salary_form": s(ws[f"I{r}"].value),
            "status": "active" if "đang làm" in status_txt else "inactive",
        }
        ctv_salaries.append({
            "employee_code": code, "month": MONTH, "company": company,
            "salary_monthly": g("K"), "salary_hourly": g("L"),
            "standard_days": g("M"), "work_days": g("N"), "work_hours": g("O"),
            "total_income": g("P"), "other_plus": g("Q"), "advance_paid": g("R"),
            "pre_tax_income": g("S"), "personal_income_tax": g("AA"),
            "net_received": g("AB"), "company_cost": g("Z"),
            "note": s(ws[f"Y{r}"].value),
        })

# ---------- 2. Chấm công TTS/CTV (BCC TTS, CTV) ----------
tts_attendance, tts_summary = [], []
if "BCC TTS, CTV" in wb.sheetnames:
    ws = wb["BCC TTS, CTV"]
    day_cols = {}
    for c in range(column_index_from_string("I"), column_index_from_string("AN") + 1):
        v = ws.cell(5, c).value
        if isinstance(v, (datetime.datetime, datetime.date)):
            day_cols[c] = v.strftime("%Y-%m-%d")
    for r in range(7, ws.max_row + 1):
        code, name = s(ws.cell(r, 2).value), s(ws.cell(r, 3).value)
        if not is_code(code) or not name: continue
        htype = norm(ws.cell(r, 8).value)  # Phân loại hợp đồng
        etype = "intern" if ("tts" in htype or "thực tập" in htype) else "ctv"
        if code not in ctv_employees:
            ctv_employees[code] = {
                "employee_code": code, "full_name": name,
                "position": s(ws.cell(r, 4).value), "department": s(ws.cell(r, 5).value),
                "company": None, "employee_type": etype,
                "join_date": d(ws.cell(r, 6).value), "contract_start_date": d(ws.cell(r, 7).value),
                "contract_type": s(ws.cell(r, 8).value) or ("Thực tập sinh" if etype == "intern" else "CTV"),
                "salary_form": None, "status": "active",
            }
        else:
            ctv_employees[code]["employee_type"] = ctv_employees[code]["employee_type"] or etype
        days = {day: s(ws.cell(r, c).value) for c, day in day_cols.items() if s(ws.cell(r, c).value)}
        if days:
            tts_attendance.append({"employee_code": code, "month": MONTH, "days": days})
        total = num(ws[f"AO{r}"].value) or 0
        tts_summary.append({
            "employee_code": code, "month": MONTH,
            "total_work_days": total, "total_leave_days": 0, "unpaid_leave_days": 0,
            "late_days": 0, "meal_days": 0, "total_paid_days": total,
        })

# ---------- 3. Biến động lương (BĐ) ----------
salary_changes = []
if "BĐ" in wb.sheetnames:
    ws = wb["BĐ"]
    TYPE_COLS = {6: "CBNV mới", 7: "Điều chỉnh chức danh", 8: "Điều chỉnh lương",
                 9: "Vào chính thức", 10: "Bổ sung/giảm trừ"}
    cur_month = None
    for r in range(4, ws.max_row + 1):
        m = s(ws.cell(r, 1).value)
        if m: cur_month = m
        name, code = s(ws.cell(r, 2).value), s(ws.cell(r, 3).value)
        if not name and not code: continue
        types = [label for c, label in TYPE_COLS.items() if s(ws.cell(r, c).value)]
        salary_changes.append({
            "month": cur_month, "employee_code": code, "full_name": name,
            "position": s(ws.cell(r, 4).value), "department": s(ws.cell(r, 5).value),
            "change_type": ", ".join(types) if types else None,
            "note": s(ws.cell(r, 11).value),
            "effective_date": d(ws.cell(r, 12).value),
        })

# ---------- 4. Truy thu chi tiết ----------
arrears = []
if "Truy thu" in wb.sheetnames:
    ws = wb["Truy thu"]
    for r in range(4, ws.max_row + 1):
        code, name = s(ws.cell(r, 2).value), s(ws.cell(r, 3).value)
        if not is_code(code) or not name: continue
        arrears.append({
            "month": MONTH, "employee_code": code, "full_name": name,
            "position": s(ws.cell(r, 4).value), "department": s(ws.cell(r, 5).value),
            "amount_shortage": num(ws.cell(r, 6).value) or 0,
            "amount_excess": num(ws.cell(r, 7).value) or 0,
            "note": s(ws.cell(r, 8).value),
        })

# ---------- 5. Thưởng Ebitda ----------
ebitda = []
if "Thưởng Ebitda" in wb.sheetnames:
    ws = wb["Thưởng Ebitda"]
    for r in range(3, ws.max_row + 1):
        code, name = s(ws.cell(r, 2).value), s(ws.cell(r, 3).value)
        if not is_code(code) or not name: continue
        ebitda.append({
            "month": MONTH, "employee_code": code, "full_name": name,
            "position": s(ws.cell(r, 4).value), "company": s(ws.cell(r, 5).value),
            "department": s(ws.cell(r, 6).value), "role": s(ws.cell(r, 7).value),
            "join_date": d(ws.cell(r, 8).value),
            "bonus": num(ws.cell(r, 9).value) or 0,
            "deduction": num(ws.cell(r, 10).value) or 0,
            "final_bonus": num(ws.cell(r, 11).value) or 0,
            "note": s(ws.cell(r, 13).value),
        })

# ---------- 6. Thang bảng lương + phụ cấp chuẩn ----------
grades = {}
if "PHU CAP" in wb.sheetnames:
    ws = wb["PHU CAP"]
    for r in range(7, ws.max_row + 1):
        grade, title = s(ws.cell(r, 1).value), s(ws.cell(r, 3).value)
        if not grade or not title: continue
        steps = {f"Bậc {i}": num(ws.cell(r, 3 + i).value) for i in range(1, 8) if num(ws.cell(r, 3 + i).value)}
        grades[(grade, norm(title))] = {
            "grade": grade, "group_name": s(ws.cell(r, 2).value), "title": title,
            "steps": steps,
            "allowance_phone": num(ws.cell(r, 11).value),
            "allowance_fuel": num(ws.cell(r, 12).value),
            "allowance_attendance": num(ws.cell(r, 13).value),
            "allowance_meal_per_day": None, "allowance_responsibility": None,
        }
if "Bảng PC, BHXH" in wb.sheetnames:
    ws = wb["Bảng PC, BHXH"]
    for r in range(5, ws.max_row + 1):
        grade, title = s(ws.cell(r, 1).value), s(ws.cell(r, 3).value)
        if not grade or not title: continue
        key = (grade, norm(title))
        row = grades.get(key) or {
            "grade": grade, "group_name": s(ws.cell(r, 2).value), "title": title,
            "steps": {}, "allowance_phone": None, "allowance_fuel": None,
            "allowance_attendance": None, "allowance_meal_per_day": None,
            "allowance_responsibility": None,
        }
        row["allowance_meal_per_day"] = num(ws.cell(r, 4).value)
        row["allowance_phone"] = num(ws.cell(r, 5).value) or row["allowance_phone"]
        row["allowance_fuel"] = num(ws.cell(r, 6).value) or row["allowance_fuel"]
        row["allowance_attendance"] = num(ws.cell(r, 7).value) or row["allowance_attendance"]
        row["allowance_responsibility"] = num(ws.cell(r, 8).value)
        grades[key] = row

# ---------- 7. KPI chi tiết (generic theo header) ----------
kpi_details = []
KPI_SHEETS = ["KPI", "KPI TVBH", "KPI CSKH", "KPI Ecom", "KPI PERFORM+ TT + TK",
              "Kho HN", "Kho HCM", "Kho Đà Nẵng", "Kho Phú Yên"]
for sheet in KPI_SHEETS:
    if sheet not in wb.sheetnames: continue
    ws = wb[sheet]
    # tìm dòng header chứa "mã n"
    hr, code_col = None, None
    for r in range(1, 8):
        for c in range(1, min(ws.max_column, 10) + 1):
            if "mã n" in norm(ws.cell(r, c).value):
                hr, code_col = r, c
                break
        if hr: break
    if not hr: continue
    headers = {}
    for c in range(1, ws.max_column + 1):
        h = s(ws.cell(hr, c).value) or s(ws.cell(hr + 1, c).value)
        if h: headers[c] = re.sub(r"\s+", " ", h)[:60]
    name_col = next((c for c, h in headers.items() if c != code_col and ("họ" in h.lower() or "tên" in h.lower())), None)
    n = 0
    for r in range(hr + 1, ws.max_row + 1):
        code = s(ws.cell(r, code_col).value)
        if not is_code(code): continue
        data = {}
        for c, h in headers.items():
            if c in (code_col, name_col): continue
            v = ws.cell(r, c).value
            if v is None or s(v) is None: continue
            data[h] = round(float(v), 2) if isinstance(v, (int, float)) else s(v)[:120]
        kpi_details.append({
            "month": MONTH, "sheet": sheet.strip(), "employee_code": code,
            "full_name": s(ws.cell(r, name_col).value) if name_col else None,
            "data": data,
        })
        n += 1
    print(f"  KPI [{sheet}]: {n} dong")

def dump(name, obj):
    with open(os.path.join(OUT, name), "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=1)

dump("seed_ctv_employees.json", list(ctv_employees.values()))
dump("seed_ctv_salaries.json", ctv_salaries)
dump("seed_ctv_attendance.json", tts_attendance)
dump("seed_ctv_attendance_summary.json", tts_summary)
dump("seed_salary_changes.json", salary_changes)
dump("seed_arrears.json", arrears)
dump("seed_ebitda.json", ebitda)
dump("seed_grades.json", list(grades.values()))
dump("seed_kpi_details.json", kpi_details)
print(f"\nCTV/TTS={len(ctv_employees)} luong_ctv={len(ctv_salaries)} cham_cong_tts={len(tts_attendance)}")
print(f"bien_dong={len(salary_changes)} truy_thu={len(arrears)} ebitda={len(ebitda)} thang_luong={len(grades)} kpi_chi_tiet={len(kpi_details)}")
