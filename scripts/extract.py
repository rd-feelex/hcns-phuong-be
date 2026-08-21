# -*- coding: utf-8 -*-
"""
Trích xuất dữ liệu từ workbook lương (mọi bố cục sheet) -> JSON seed.
Cách dùng:  python scripts/extract.py "duong-dan-file.xlsx" [password]
Nhận diện cột theo TIÊU ĐỀ (header) nên chịu được việc các sheet có bố cục khác nhau.
"""
import msoffcrypto, io, json, datetime, sys, os, re
import openpyxl
from openpyxl.utils import get_column_letter

SRC = sys.argv[1]
PASSWORD = sys.argv[2] if len(sys.argv) > 2 else "2019OZ"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")
MONTH = "2026-07"

# Thứ tự ưu tiên xác định "công ty chính" của nhân viên xuất hiện ở nhiều sheet
PRIORITY = ["OZO", "XTRADING", "XCentre", "XExtra", "XGlobal", "VTM", "Feelex", "X Group", "ENDLESS"]
SHEET_COMPANY = {
    "OZO": "OZO", "XTRADING": "XTRADING", "FEELEX": "Feelex", "X GROUP": "X Group",
    "ENDLESS": "ENDLESS", "VTM": "VTM", "XCENTRE": "XCentre", "X EXTRA": "XExtra", "X GLOBAL": "XGlobal",
}

def norm(v):
    if v is None: return ""
    return re.sub(r"\s+", " ", str(v)).strip().lower()

def num(v): return round(float(v), 2) if isinstance(v, (int, float)) else None
def s(v):
    t = str(v).strip() if v is not None else None
    return t if t else None
def d(v): return v.strftime("%Y-%m-%d") if isinstance(v, (datetime.datetime, datetime.date)) else None

with open(SRC, "rb") as f:
    of = msoffcrypto.OfficeFile(f)
    of.load_key(password=PASSWORD)
    dec = io.BytesIO()
    of.decrypt(dec)
dec.seek(0)
wb = openpyxl.load_workbook(dec, data_only=True)

class Sheet:
    """Bảng lương một công ty với cột được nhận diện theo header."""
    def __init__(self, ws):
        self.ws = ws
        # Tìm dòng header: cột B chứa "mã nv"/"mã nhân viên"
        self.hr = None
        for r in range(1, 12):
            b = norm(ws.cell(r, 2).value)
            if "mã nv" in b or "mã nhân viên" in b:
                self.hr = r
                break
        if self.hr is None:
            raise ValueError(f"Không tìm thấy dòng header trong sheet {ws.title}")
        # group: ô ở dòng header; sub: ô ở dòng dưới. Ô group kéo dài tới group kế tiếp.
        self.group, self.sub = {}, {}
        cur_group = ""
        for c in range(1, ws.max_column + 1):
            g = norm(ws.cell(self.hr, c).value)
            if g: cur_group = g
            self.group[c] = cur_group
            self.sub[c] = norm(ws.cell(self.hr + 1, c).value)

    def cols_matching(self, cand):
        """Các cột có sub hoặc group chứa chuỗi cand."""
        out = []
        for c in range(1, self.ws.max_column + 1):
            if cand in self.sub[c] or cand in self.group[c]:
                out.append(c)
        return out

    def find(self, *cands, after=None, before=None, exclude=""):
        # Ưu tiên khớp ở dòng tiêu đề con (sub) trước — tiêu đề nhóm dài dễ chứa
        # từ khóa gây match nhầm (VD "Tổng lương ... phụ cấp ăn trưa").
        for source in ("sub", "group"):
            for cand in cands:
                for c in range(1, self.ws.max_column + 1):
                    txt = self.sub[c] if source == "sub" else self.group[c]
                    if cand not in txt: continue
                    label = self.sub[c] + " " + self.group[c]
                    if exclude and exclude in label: continue
                    if after is not None and c <= after: continue
                    if before is not None and c >= before: continue
                    return c
        return None

    def find_last(self, cand, before):
        cols = [c for c in self.cols_matching(cand) if c < before]
        return cols[-1] if cols else None

    def val(self, r, c):
        return self.ws.cell(r, c).value if c else None

    def data_rows(self):
        for r in range(self.hr + 2, self.ws.max_row + 1):
            stt = self.ws.cell(r, 1).value
            code = s(self.ws.cell(r, 2).value)
            name = s(self.ws.cell(r, 3).value)
            if not isinstance(stt, (int, float)): continue
            if not code or not name: continue
            if len(code) < 6 or code.isdigit(): continue  # dòng chú thích/đánh số
            yield r, code, name

employees, salaries, deps, emails = {}, [], {}, {}

def parse_company_sheet(ws, company):
    sh = Sheet(ws)
    F = {}
    F["position"] = sh.find("chức danh")
    F["department"] = sh.find("phòng ban", "bộ phận")
    F["contract_start"] = sh.find("bắt đầu hđlđ", "ngày chính thức")
    F["join"] = sh.find("ngày vào cty", "ngày bắt đầu", exclude="hđlđ")
    F["probation_end"] = sh.find("kết thúc thử việc")
    F["contract_type"] = sh.find("loại hình hợp đồng", "phân loại hợp đồng", "loại hđ")
    F["grade"] = sh.find("ngạch lương", "cấp bậc")
    F["form"] = sh.find("hình thức lương")
    F["agreed"] = sh.find("thỏa thuận", "thu nhập chính thức", "tổng lương")
    F["base"] = sh.find("lương cơ bản", exclude="đóng")
    F["insurance"] = sh.find("mức lương cơ bản ( đóng", "lương đóng bhxh")
    F["pc_phone"] = sh.find("phụ cấp điện thoại")
    F["pc_fuel"] = sh.find("phụ cấp xăng xe")
    F["pc_att"] = sh.find("phụ cấp chuyên cần")
    F["pc_resp"] = sh.find("phụ cấp trách nhiệm")
    F["pc_meal"] = sh.find("phụ cấp ăn ca/ ngày", "phụ cấp ăn trưa")
    F["std_days"] = sh.find("công chuẩn")
    kpi_pct = sh.find("% kpi thực đạt")
    F["kpi_pct"] = kpi_pct
    F["kpi_target"] = sh.find("mức thưởng") or (kpi_pct and sh.find_last("lương đánh giá hiệu q", before=kpi_pct))
    F["kpi_actual"] = kpi_pct and sh.find("lương đánh giá hiệu q", after=kpi_pct)
    F["biz_bonus"] = sh.find("thưởng doanh thu")
    F["inc_bonus"] = sh.find("thưởng khuyến khích")
    F["other_ded"] = sh.find("giảm trừ khác")
    F["total_bonus"] = sh.find("tổng thưởng")
    F["arrears"] = sh.find("truy thu/ hỗ trợ khác", "điều chỉnh tăng giảm")
    F["arrears_note"] = sh.find("nội dung truy thu", "nội dung (+)")
    F["total_income"] = sh.find("tổng thu nhập")
    F["dependents"] = sh.find("số người")
    F["si_emp"] = sh.find("tổng người ld chịu", "tổng người lđ chịu", "total 10.5")
    F["si_total"] = sh.find("tổng nộp bhxh", "tổng bhxh")
    F["union_emp"] = sh.find("nld chịu", "nlđ chịu")
    F["tax"] = sh.find("thuế tncn")
    F["net"] = sh.find("nlđ thực lĩnh", "thu nhập net")
    F["advance"] = sh.find("đã tạm ứng", "tạm ứng/ truy thu", "tạm ứng")
    F["final"] = sh.find("thực nhận", "lương thực lĩnh")
    F["cost"] = sh.find("chi phí công ty", "chi phí cty")
    F["email"] = sh.find("mail")

    # Công tính lương: ưu tiên tên trực tiếp, fallback sub "tổng" của nhóm "công"
    F["paid_days"] = sh.find("công tính lương", "công được tính lương", exclude="ăn ca")
    if not F["paid_days"]:
        for c in sh.cols_matching("tổng"):
            if sh.group[c] == "công" and sh.sub[c] == "tổng":
                F["paid_days"] = c
                break
    # Lương theo ngày công: nhóm có sub "tổng" thì lấy sub, không thì lấy chính cột nhóm
    lw = sh.cols_matching("lương theo ngày công")
    F["salary_days"] = None
    if lw:
        subs = [c for c in lw if sh.sub[c] == "tổng"]
        F["salary_days"] = subs[0] if subs else lw[0]
    F["total_pc"] = sh.find("tổng pc")

    n_rows = 0
    for r, code, name in sh.data_rows():
        n_rows += 1
        g = lambda k: num(sh.val(r, F.get(k)))
        if code not in employees:
            employees[code] = {
                "employee_code": code, "full_name": name,
                "position": s(sh.val(r, F["position"])),
                "department": s(sh.val(r, F["department"])),
                "company": company,
                "join_date": d(sh.val(r, F["join"])),
                "probation_end_date": d(sh.val(r, F["probation_end"])),
                "contract_start_date": d(sh.val(r, F["contract_start"])),
                "contract_type": s(sh.val(r, F["contract_type"])),
                "salary_grade": s(sh.val(r, F["grade"])),
                "salary_form": s(sh.val(r, F["form"])),
                "status": "active",
            }
        dep = sh.val(r, F["dependents"])
        if code not in deps and isinstance(dep, (int, float)):
            deps[code] = int(dep)
        em = s(sh.val(r, F["email"]))
        if em and "@" in em and code not in emails:
            emails[code] = em
        meal = g("pc_meal")
        if meal and meal > 200000:  # sheet ghi theo tháng -> quy về /ngày
            std = g("std_days") or 26
            meal = round(meal / std, 0)
        salaries.append({
            "employee_code": code, "month": MONTH, "company": company,
            "agreed_salary": g("agreed"), "base_salary": g("base"),
            "insurance_salary": g("insurance"),
            "allowance_phone": g("pc_phone"), "allowance_fuel": g("pc_fuel"),
            "allowance_attendance": g("pc_att"), "allowance_responsibility": g("pc_resp"),
            "allowance_meal_per_day": meal,
            "kpi_salary_target": g("kpi_target"),
            "standard_workdays": g("std_days"), "paid_workdays": g("paid_days"),
            "salary_by_workdays": g("salary_days"), "total_allowance": g("total_pc"),
            "kpi_percent": g("kpi_pct"), "kpi_salary_actual": g("kpi_actual"),
            "business_bonus": g("biz_bonus"), "incentive_bonus": g("inc_bonus"),
            "other_deduction": g("other_ded"), "total_bonus": g("total_bonus"),
            "arrears": g("arrears"), "arrears_note": s(sh.val(r, F["arrears_note"])),
            "total_income": g("total_income"),
            "social_insurance_total": g("si_total"), "social_insurance_employee": g("si_emp"),
            "union_fee_employee": g("union_emp"),
            "personal_income_tax": g("tax"), "net_income": g("net"),
            "advance_paid": g("advance"), "final_received": g("final"),
            "company_cost": g("cost"),
        })
    total_final = sum((x["final_received"] or 0) for x in salaries if x["company"] == company)
    print(f"  {company:10s} header=R{sh.hr} rows={n_rows:3d} tong_thuc_nhan={total_final:,.0f} "
          f"(final=cot {get_column_letter(F['final']) if F['final'] else '?'})")

# Sắp sheet theo thứ tự ưu tiên công ty chính
sheet_by_company = {}
for name in wb.sheetnames:
    key = name.strip().upper()
    if key in SHEET_COMPANY:
        sheet_by_company[SHEET_COMPANY[key]] = name

print("Trích xuất bảng lương:")
for comp in PRIORITY:
    if comp in sheet_by_company:
        parse_company_sheet(wb[sheet_by_company[comp]], comp)

# ----- Chấm công BCC NV -----
from openpyxl.utils import column_index_from_string
ws = wb["BCC NV"]
day_cols = {}
for c in range(column_index_from_string("K"), column_index_from_string("AO") + 1):
    v = ws.cell(5, c).value
    if isinstance(v, (datetime.datetime, datetime.date)):
        day_cols[c] = v.strftime("%Y-%m-%d")

attendance, att_summary = [], []
for r in range(8, ws.max_row + 1):
    code, name = s(ws.cell(r, 2).value), s(ws.cell(r, 3).value)
    if not code or not name: continue
    if len(code) < 6 or code.isdigit(): continue
    if code not in employees:
        employees[code] = {
            "employee_code": code, "full_name": name,
            "position": s(ws.cell(r, 4).value), "department": s(ws.cell(r, 5).value),
            "company": s(ws.cell(r, 6).value),
            "join_date": d(ws.cell(r, 7).value), "probation_end_date": d(ws.cell(r, 8).value),
            "contract_start_date": d(ws.cell(r, 9).value), "contract_type": s(ws.cell(r, 10).value),
            "salary_grade": None, "salary_form": None, "status": "active",
        }
    days = {day: s(ws.cell(r, c).value) for c, day in day_cols.items() if s(ws.cell(r, c).value)}
    if days:
        attendance.append({"employee_code": code, "month": MONTH, "days": days})
    col = lambda L: num(ws.cell(r, column_index_from_string(L)).value)
    att_summary.append({
        "employee_code": code, "month": MONTH,
        "total_work_days": col("AQ") or 0, "total_leave_days": col("AS") or 0,
        "unpaid_leave_days": col("BB") or 0, "late_days": col("BC") or 0,
        "meal_days": col("BD") or 0, "total_paid_days": col("BG") or 0,
    })

for e in employees.values():
    e["num_dependents"] = deps.get(e["employee_code"], 0)
    if emails.get(e["employee_code"]):
        e["email"] = emails[e["employee_code"]]

primary = {e["employee_code"]: e["company"] for e in employees.values()}
inputs, seen = [], set()
for srow in salaries:
    code = srow["employee_code"]
    if srow["company"] != primary.get(code) or code in seen: continue
    seen.add(code)
    inputs.append({
        "employee_code": code, "month": srow["month"],
        "kpi_percent": srow["kpi_percent"] if srow["kpi_percent"] is not None else 1,
        "business_bonus": srow["business_bonus"] or 0,
        "incentive_bonus": srow["incentive_bonus"] or 0,
        "other_deduction": srow["other_deduction"] or 0,
        "arrears": srow["arrears"] or 0,
        "arrears_note": srow["arrears_note"] if srow["arrears_note"] not in (None, "0") else None,
        "advance_paid": srow["advance_paid"] or 0,
    })

def dump(name, obj):
    with open(os.path.join(OUT, name), "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=1)

dump("seed_employees.json", list(employees.values()))
dump("seed_salaries.json", salaries)
dump("seed_attendance.json", attendance)
dump("seed_attendance_summary.json", att_summary)
dump("seed_payroll_inputs.json", inputs)
grand = sum((x["final_received"] or 0) for x in salaries)
print(f"\nTONG: employees={len(employees)} salaries={len(salaries)} attendance={len(attendance)} "
      f"inputs={len(inputs)} | tong_thuc_nhan toan bo = {grand:,.0f}")
